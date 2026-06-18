"""
Generate defense presentation PPTX from the HCMUT template.
Uses the existing slide master/layouts and HCMUT branding from the template.
"""

import os
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

TEMPLATE = "academic-paper-vietnamese/template.pptx"
OUTPUT = "academic-paper-vietnamese/defense_presentation_v2.pptx"
FIGURES = "academic-paper-vietnamese/figures"

# Colors matching template style
DARK_BLUE = RGBColor(0x00, 0x32, 0x66)
ACCENT_BLUE = RGBColor(0x00, 0x70, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
GREEN = RGBColor(0x00, 0x80, 0x00)
RED = RGBColor(0xCC, 0x00, 0x00)


def set_text(shape, text, font_size=18, bold=False, color=None, alignment=None):
    """Set text on a shape, clearing existing content."""
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    if color:
        p.font.color.rgb = color
    if alignment:
        p.alignment = alignment


def add_bullet_text(text_frame, items, font_size=16, bold_prefix=False, color=None):
    """Add bullet items to a text frame. Each item can be 'bold: rest' format."""
    text_frame.clear()
    for i, item in enumerate(items):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()

        p.space_after = Pt(4)
        p.space_before = Pt(2)
        p.level = 0

        if bold_prefix and ": " in item:
            bold_part, rest = item.split(": ", 1)
            run1 = p.add_run()
            run1.text = bold_part + ": "
            run1.font.size = Pt(font_size)
            run1.font.bold = True
            if color:
                run1.font.color.rgb = color
            run2 = p.add_run()
            run2.text = rest
            run2.font.size = Pt(font_size)
            if color:
                run2.font.color.rgb = color
        else:
            run = p.add_run()
            run.text = item
            run.font.size = Pt(font_size)
            if color:
                run.font.color.rgb = color


def add_section_label(slide, text, prs):
    """Add a section label textbox in top-right area (matching template style)."""
    from pptx.util import Inches, Pt
    txBox = slide.shapes.add_textbox(
        Inches(9.8), Inches(0.2), Inches(3.2), Inches(0.45)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = ACCENT_BLUE


def add_image_centered(slide, img_path, left, top, max_width, max_height):
    """Add image with aspect ratio preserved, centered in the given area."""
    if not os.path.exists(img_path):
        print(f"  WARNING: Image not found: {img_path}")
        return None
    pic = slide.shapes.add_picture(img_path, left, top, max_width)
    # Scale to fit within max_height
    if pic.height > max_height:
        ratio = max_height / pic.height
        pic.width = int(pic.width * ratio)
        pic.height = max_height
    # Also check width
    if pic.width > max_width:
        ratio = max_width / pic.width
        pic.height = int(pic.height * ratio)
        pic.width = max_width
    # Center horizontally within area
    pic.left = left + (max_width - pic.width) // 2
    pic.top = top + (max_height - pic.height) // 2
    return pic


def build_presentation():
    prs = Presentation(TEMPLATE)

    # Delete all existing slides (keep only master/layouts)
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

    title_layout = prs.slide_layouts[0]       # "Title Slide"
    content_layout = prs.slide_layouts[1]     # "Title and Content"
    two_content_layout = prs.slide_layouts[3] # "Two Content"
    section_layout = prs.slide_layouts[2]     # "Section Header"
    title_only_layout = prs.slide_layouts[5]  # "Title Only"
    blank_layout = prs.slide_layouts[6]       # "Blank"

    # ============================================================
    # SLIDE 1 — Title
    # ============================================================
    slide = prs.slides.add_slide(title_layout)
    slide.placeholders[0].text = (
        "Xây dựng framework tạo dựng mô hình AI\n"
        "cho các ứng dụng theo dõi chuyển động của con người"
    )
    for run in slide.placeholders[0].text_frame.paragraphs[0].runs:
        run.font.size = Pt(28)
    slide.placeholders[1].text = (
        "\nGVHD: TS. Lê Trọng Nhân\n"
        "Học viên: Nguyễn Trương Minh Hoàng - 2270757"
    )
    for para in slide.placeholders[1].text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(16)
    print("Slide 1: Title")

    # ============================================================
    # SLIDE 2 — Outline
    # ============================================================
    slide = prs.slides.add_slide(content_layout)
    slide.placeholders[0].text = "Nội dung trình bày"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    sections = [
        "1. Bối cảnh & Vấn đề nghiên cứu",
        "2. Mục tiêu & Đóng góp",
        "3. Kiến trúc & Tính năng cốt lõi",
        "4. Kết quả thực nghiệm",
        "5. So sánh với nền tảng thương mại",
        "6. Kết luận & Hướng phát triển",
    ]
    for i, sec in enumerate(sections):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = sec
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        p.space_after = Pt(12)
    print("Slide 2: Outline")

    # ============================================================
    # SLIDE 3 — Background & Problem
    # ============================================================
    slide = prs.slides.add_slide(two_content_layout)
    slide.placeholders[0].text = "Bối cảnh & Vấn đề nghiên cứu"
    add_section_label(slide, "Bối cảnh", prs)

    add_bullet_text(slide.placeholders[1].text_frame, [
        "HAR (Human Activity Recognition): nhận dạng hoạt động từ cảm biến quán tính (IMU 6 trục)",
        "Ứng dụng: y tế (phát hiện té ngã), thể thao, công nghiệp",
        "Yêu cầu: xử lý ngay trên thiết bị (edge), không phụ thuộc mạng",
        "Khoảng cách triển khai: hầu hết nghiên cứu dừng ở Python prototype",
        "Nền tảng thương mại (Edge Impulse, SensiML): hộp đen, đắt tiền",
    ], font_size=15, bold_prefix=True)

    add_bullet_text(slide.placeholders[2].text_frame, [
        "5 rào cản chính:",
        "  • Pipeline liên ngành phức tạp",
        "  • Parity gap: Python ≠ C++",
        "  • Thiếu công cụ mã nguồn mở end-to-end",
        "  • Thiếu hỗ trợ ML cổ điển (RF, SVM)",
        "  • Nền tảng thương mại: hộp đen, khó kiểm chứng",
    ], font_size=15)
    print("Slide 3: Background & Problem")

    # ============================================================
    # SLIDE 4 — Objectives & Contributions
    # ============================================================
    slide = prs.slides.add_slide(two_content_layout)
    slide.placeholders[0].text = "Mục tiêu & Đóng góp"
    add_section_label(slide, "Mục tiêu", prs)

    add_bullet_text(slide.placeholders[1].text_frame, [
        "Mục tiêu nghiên cứu:",
        "  ✓ Framework end-to-end: dữ liệu → mã C++ trên MCU",
        "  ✓ Giao diện tương tác, không cần kỹ năng nhúng",
        "  ✓ Đảm bảo tương đồng huấn luyện–triển khai",
        "  ✓ Hỗ trợ đa nền tảng, đa thuật toán",
        "  ✓ Xác thực trên thiết bị thực (XIAO nRF52840)",
    ], font_size=15)

    add_bullet_text(slide.placeholders[2].text_frame, [
        "Đóng góp chính:",
        "  ✓ Phân đoạn cửa sổ kéo thả tương tác",
        "  ✓ Pipeline đảm bảo parity Python ↔ C++",
        "  ✓ Chiến lược đệm cửa sổ có cơ sở vật lý",
        "  ✓ 6 chế độ đặc trưng (orientation-invariant)",
        "  ✓ Tạo mã đa nền tảng: Arduino, ARM, Zephyr, MicroPython",
    ], font_size=15)
    print("Slide 4: Objectives & Contributions")

    # ============================================================
    # SLIDE 5 — Architecture (6-tab pipeline with screenshots)
    # ============================================================
    slide = prs.slides.add_slide(title_only_layout)
    slide.placeholders[0].text = "Kiến trúc Framework: 6 Tab Tuần tự"
    add_section_label(slide, "Kiến trúc", prs)

    # Add 6 tab boxes as a horizontal pipeline using text boxes
    tab_names = ["Dữ liệu", "Tiền xử lý", "Đặc trưng", "Huấn luyện", "Tạo mã", "Kiểm tra"]
    tab_descs = ["CSV\n100 Hz", "Lọc\nKéo thả", "63 feat\n6 chế độ", "RF/SVM\nNN/CNN", "C++/MPy\nĐa nền tảng", "Serial\nXIAO"]
    start_x = Inches(0.8)
    tab_w = Inches(1.7)
    tab_h = Inches(0.8)
    gap = Inches(0.3)
    y_tab = Inches(1.8)

    for i, (name, desc) in enumerate(zip(tab_names, tab_descs)):
        x = start_x + i * (tab_w + gap)
        # Tab box
        txBox = slide.shapes.add_textbox(x, y_tab, tab_w, tab_h)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"Tab {i+1}: {name}"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = WHITE
        # Add fill
        from pptx.oxml.ns import qn
        spPr = txBox._element.find(qn("a:spPr") if txBox._element.find(qn("a:spPr")) is not None else ".//{http://schemas.openxmlformats.org/drawingml/2006/main}spPr")
        if spPr is None:
            spPr = etree.SubElement(txBox._element.find(qn("p:txBody")).getparent(), qn("p:spPr"))
        # Use the shape's spPr
        sp_el = txBox._element
        spPr = sp_el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}spPr")
        if spPr is None:
            spPr = etree.SubElement(sp_el, "{http://schemas.openxmlformats.org/drawingml/2006/main}spPr")
        solidFill = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
        srgbClr = etree.SubElement(solidFill, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", val="003266")

        # Description below
        txBox2 = slide.shapes.add_textbox(x, y_tab + tab_h + Inches(0.05), tab_w, Inches(0.5))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = desc
        run2.font.size = Pt(10)
        run2.font.color.rgb = DARK_GRAY

        # Arrow between tabs
        if i < 5:
            arrow_x = x + tab_w
            arrow_y = y_tab + tab_h // 2
            arrow = slide.shapes.add_textbox(arrow_x, arrow_y - Inches(0.1), gap, Inches(0.3))
            atf = arrow.text_frame
            ap = atf.paragraphs[0]
            ap.alignment = PP_ALIGN.CENTER
            ar = ap.add_run()
            ar.text = "→"
            ar.font.size = Pt(16)
            ar.font.bold = True
            ar.font.color.rgb = ACCENT_BLUE

    # Technology stack at bottom
    tech_items = [
        ("Công nghệ:", "Python 3 · Dash 2.14 · scikit-learn · PyTorch"),
        ("Lưu trữ:", "persistent_data/ (CSV → joblib → .h/.ino) — không cần DB"),
        ("Thiết bị:", "Seeed XIAO nRF52840 — ARM Cortex-M4F 64 MHz — IMU LSM6DS3"),
    ]
    y_tech = Inches(3.5)
    for j, (label, value) in enumerate(tech_items):
        txBox = slide.shapes.add_textbox(Inches(1.0), y_tech + j * Inches(0.45), Inches(11.0), Inches(0.4))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        run1 = p.add_run()
        run1.text = label + " "
        run1.font.size = Pt(14)
        run1.font.bold = True
        run1.font.color.rgb = DARK_BLUE
        run2 = p.add_run()
        run2.text = value
        run2.font.size = Pt(14)
        run2.font.color.rgb = DARK_GRAY
    print("Slide 5: Architecture")

    # ============================================================
    # SLIDE 6 — Preprocessing (with screenshot)
    # ============================================================
    slide = prs.slides.add_slide(title_only_layout)
    slide.placeholders[0].text = "Tiền xử lý Tương tác"
    add_section_label(slide, "Tính năng cốt lõi", prs)

    # Left: key points
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(4.5), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    add_bullet_text(tf, [
        "Phân đoạn cửa sổ kéo thả (novel):",
        "  • Biểu đồ tín hiệu tương tác Plotly",
        "  • Kéo thả vùng hoạt động trên đồ thị",
        "  • Giảm ~60–80% thời gian tiền xử lý",
        "",
        "Bộ lọc: Butterworth / Savitzky-Golay / Kalman",
        "",
        "Đệm cửa sổ: edge-value replication",
        "  • Bảo toàn đặc tính vật lý của tín hiệu",
        "  • Tránh giá trị acc_mag = 0 (bất khả thi)",
    ], font_size=14)

    # Right: screenshot
    img = os.path.join(FIGURES, "ui_preprocessing_draggable.png")
    add_image_centered(slide, img, Inches(5.2), Inches(1.5), Inches(7.5), Inches(5.5))
    print("Slide 6: Preprocessing")

    # ============================================================
    # SLIDE 7 — Feature Extraction (with screenshot)
    # ============================================================
    slide = prs.slides.add_slide(title_only_layout)
    slide.placeholders[0].text = "Trích xuất Đặc trưng & Parity"
    add_section_label(slide, "Tính năng cốt lõi", prs)

    # Left: feature modes
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(5.5), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    add_bullet_text(tf, [
        "6 chế độ đặc trưng:",
        "  • orientation_invariant (63) ← mặc định",
        "  • time_domain (90)",
        "  • all (156) | frequency_domain (66)",
        "  • orientation_invariant_time_only (33) | raw (6)",
        "",
        "15 thống kê/tín hiệu: mean, std, min, max, range, median, Q25, Q75, IQR, skewness, kurtosis, RMS, energy, ZCR, MCR",
        "",
        "Đảm bảo parity Python ↔ C++:",
        "  • Cùng công thức, cùng thứ tự, cùng đơn vị",
        "  • Reorder tham số tại thời điểm sinh mã",
        "  • Chuẩn hóa 1 lần duy nhất trong pipeline",
    ], font_size=13)

    # Right: screenshot
    img = os.path.join(FIGURES, "ui_feature_engineering.png")
    add_image_centered(slide, img, Inches(6.3), Inches(1.5), Inches(6.5), Inches(5.5))
    print("Slide 7: Feature Extraction")

    # ============================================================
    # SLIDE 8 — Code Generation (with screenshot)
    # ============================================================
    slide = prs.slides.add_slide(title_only_layout)
    slide.placeholders[0].text = "Tạo Mã Triển khai Đa Nền tảng"
    add_section_label(slide, "Tính năng cốt lõi", prs)

    # Left: code gen info
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(5.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    add_bullet_text(tf, [
        "Nền tảng đích:",
        "  • Arduino (XIAO, ESP32, M5Stack, Teensy)",
        "  • ARM Cortex-M generic (C thuần)",
        "  • ESP-IDF, Zephyr RTOS",
        "  • MicroPython (không cần NumPy)",
        "  • TFLite Micro / ONNX Runtime (NN/CNN)",
        "",
        "4 chế độ tối ưu:",
        "  accuracy | balanced | speed | power",
        "",
        "Kiến trúc: BaseCodeGenerator → subclass theo từng loại mô hình (NN, RF, SVM, CNN)",
    ], font_size=14)

    # Right: screenshot
    img = os.path.join(FIGURES, "ui_code_generation.png")
    add_image_centered(slide, img, Inches(5.5), Inches(1.5), Inches(7.3), Inches(5.5))
    print("Slide 8: Code Generation")

    # ============================================================
    # SLIDE 9 — Training Results (with confusion matrices)
    # ============================================================
    slide = prs.slides.add_slide(title_only_layout)
    slide.placeholders[0].text = "Kết quả Huấn luyện"
    add_section_label(slide, "Kết quả", prs)

    # Results table as text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.5), Inches(4.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    # 3-class header
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Bài toán 3 lớp: 100% tất cả thuật toán"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = GREEN

    p = tf.add_paragraph()
    run = p.add_run()
    run.text = ""
    run.font.size = Pt(6)

    # 5-class header
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = "Bài toán 5 lớp (229 mẫu kiểm tra):"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE

    results_5class = [
        ("PyTorch MLP", "99,13%", "F1=0,990"),
        ("PyTorch 1D-CNN", "99,13%", "F1=0,992"),
        ("Neural Network (sklearn)", "98,69%", "F1=0,984"),
        ("Random Forest", "98,25%", "F1=0,981"),
        ("SVM (RBF)", "97,38%", "F1=0,972"),
    ]
    for name, acc, f1 in results_5class:
        p = tf.add_paragraph()
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = f"  {name}: {acc}  ({f1})"
        run.font.size = Pt(14)

    p = tf.add_paragraph()
    p.space_before = Pt(10)
    run = p.add_run()
    run.text = "→ Chênh lệch chỉ 1,75% → chất lượng đặc trưng quan trọng hơn lựa chọn thuật toán"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = ACCENT_BLUE

    # Right: confusion matrix images
    img1 = os.path.join(FIGURES, "ui_mlp_training_c5_confusion.png")
    add_image_centered(slide, img1, Inches(6.3), Inches(1.4), Inches(6.5), Inches(5.5))
    print("Slide 9: Training Results")

    # ============================================================
    # SLIDE 10 — Device Deployment
    # ============================================================
    slide = prs.slides.add_slide(two_content_layout)
    slide.placeholders[0].text = "Kết quả Triển khai Thiết bị"
    add_section_label(slide, "Kết quả", prs)

    add_bullet_text(slide.placeholders[1].text_frame, [
        "Thiết bị: Seeed XIAO nRF52840 (Cortex-M4F 64MHz)",
        "",
        "Hiệu suất trên thiết bị:",
        "  ✓ 1D-CNN: ổn định cả 3 và 5 lớp",
        "  • NN sklearn: lỗi ở một số điều kiện",
        "  • RF: ổn định cho 3 lớp",
        "",
        "Tài nguyên (Neural Network):",
        "  • Flash: ~95 KB (trọng số + FE + scaler)",
        "  • SRAM: < 8 KB",
        "  • Suy luận: < 100 ms",
    ], font_size=15, bold_prefix=True)

    add_bullet_text(slide.placeholders[2].text_frame, [
        "So sánh với TFLite Micro:",
        "                 Framework    TFLite",
        "  Flash (NN):   95 KB         ~245 KB",
        "  RF/SVM:        ✓ Hỗ trợ     ✗ Không",
        "  Mã đọc được:  ✓               ✗",
        "",
        "Bài học: chênh lệch 1,75% offline không dự đoán được hiệu suất thực trên thiết bị",
        "",
        "→ Tiêu chí chọn mô hình phải bao gồm tính ổn định triển khai",
    ], font_size=14)
    print("Slide 10: Device Deployment")

    # ============================================================
    # SLIDE 11 — Comparison with Commercial Platforms
    # ============================================================
    slide = prs.slides.add_slide(content_layout)
    slide.placeholders[0].text = "So sánh với Nền tảng Thương mại"
    add_section_label(slide, "So sánh", prs)

    tf = slide.placeholders[1].text_frame
    tf.clear()

    # Build comparison as formatted text
    headers = f"{'Tiêu chí':<30} {'Edge Impulse':<18} {'SensiML':<18} {'Framework này':<18}"
    rows = [
        ("Chi phí", "$20–99/tháng", "$99–500/tháng", "Miễn phí"),
        ("Mã nguồn mở", "✗", "✗", "✓"),
        ("Minh bạch mã sinh ra", "Hộp đen", "Hộp đen", "✓ Toàn bộ"),
        ("Đảm bảo parity", "Không rõ", "Không rõ", "✓ Bit-exact"),
        ("Hỗ trợ RF/SVM", "Giới hạn", "Có", "✓"),
        ("F1 (3 lớp, cùng data)", "0,96", "—", "1,00"),
        ("Flash (NN)", "51%", "—", "18–23%"),
    ]

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = headers
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.name = "Consolas"
    run.font.color.rgb = DARK_BLUE

    for label, ei, sm, fw in rows:
        p = tf.add_paragraph()
        p.space_before = Pt(1)
        run = p.add_run()
        run.text = f"{label:<30} {ei:<18} {sm:<18} {fw:<18}"
        run.font.size = Pt(13)
        run.font.name = "Consolas"

    p = tf.add_paragraph()
    p.space_before = Pt(14)
    run = p.add_run()
    run.text = '→ Tính minh bạch hoàn toàn là yêu cầu kỹ thuật thiết yếu cho edge ML, không chỉ triết lý mã nguồn mở'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = ACCENT_BLUE
    print("Slide 11: Commercial Comparison")

    # ============================================================
    # SLIDE 12 — Conclusion
    # ============================================================
    slide = prs.slides.add_slide(two_content_layout)
    slide.placeholders[0].text = "Kết luận & Hướng phát triển"
    add_section_label(slide, "Kết luận", prs)


    add_bullet_text(slide.placeholders[1].text_frame, [
        "Đã đạt được:",
        "  ✓ Framework end-to-end, mã nguồn mở",
        "  ✓ 97–99% (5 lớp), 100% (3 lớp)",
        "  ✓ Tương đồng huấn luyện–triển khai đảm bảo",
        "  ✓ Triển khai thành công trên XIAO",
        "  ✓ Vượt Edge Impulse: F1 1,00 vs 0,96",
    ], font_size=16)

    add_bullet_text(slide.placeholders[2].text_frame, [
        "Hướng phát triển:",
        "  • Dữ liệu đa đối tượng, đa hướng gắn",
        "  • Học liên tục trên thiết bị",
        "  • Sensor fusion (IMU + nhiệt + nhịp tim)",
        "  • CAE học đặc trưng tự động",
        "  • Tối ưu đa mục tiêu",
        "  • Phát hiện bất thường",
    ], font_size=16)
    print("Slide 12: Conclusion")

    # ============================================================
    # SLIDE 13 — Thank You
    # ============================================================
    slide = prs.slides.add_slide(title_layout)
    slide.placeholders[0].text = "Cảm ơn Quý Thầy/Cô & Hội đồng\n\nQ&A"
    for para in slide.placeholders[0].text_frame.paragraphs:
        para.alignment = PP_ALIGN.CENTER
        for run in para.runs:
            run.font.size = Pt(32)

    slide.placeholders[1].text = (
        "\nSinh viên: Nguyễn Trương Minh Hoàng — MSSV: 2270757\n"
        "GVHD: TS. Lê Trọng Nhân\n\n"
        '"Tính minh bạch không chỉ là triết lý mã nguồn mở\n'
        '— đây là yêu cầu kỹ thuật thiết yếu cho edge AI đáng tin cậy."'
    )
    for para in slide.placeholders[1].text_frame.paragraphs:
        para.alignment = PP_ALIGN.CENTER
        for run in para.runs:
            run.font.size = Pt(14)
    print("Slide 13: Thank You")

    # ============================================================
    # Save
    # ============================================================
    prs.save(OUTPUT)
    print(f"\nSaved: {OUTPUT}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_presentation()
