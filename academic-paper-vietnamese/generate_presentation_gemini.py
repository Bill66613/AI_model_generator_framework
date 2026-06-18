from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
 
def add_slide(prs, title_text, content_text):
    slide_layout = prs.slide_layouts[1] # Bullet layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    tf = slide.placeholders[1].text_frame
    tf.text = content_text
 
def create_presentation():
    prs = Presentation()
 
    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Xây dựng Framework tạo dựng mô hình AI cho các ứng dụng theo dõi chuyển động con người"
    subtitle.text = "Học viên: Nguyễn Trương Minh Hoàng\nGVHD: TS. Lê Trọng Nhân\nTrường Đại học Bách Khoa - ĐHQG TP.HCM"
 
    # Danh sách 25 Slide
    slides_content = [
        ("Tóm tắt nghiên cứu", "- Framework toàn diện cho Edge AI trong theo dõi chuyển động.\n- Tiền xử lý tương tác kéo-thả trực quan.\n- Sinh mã triển khai đa nền tảng (Arduino, ARM, ESP32...).\n- Đảm bảo tính tương đồng (Parity) tuyệt đối."),
        ("Bối cảnh: Wearables & IMU", "- Sự phát triển của thiết bị đeo thông minh.\n- Vai trò của cảm biến IMU (Gia tốc & Con quay hồi chuyển).\n- Ưu điểm: Bảo mật, chi phí thấp, tiêu thụ năng lượng thấp."),
        ("Xu hướng: Xử lý tại biên (Edge AI)", "- Giảm độ trễ (phản hồi < 20ms).\n- Không phụ thuộc kết nối mạng Cloud.\n- Bảo mật dữ liệu thô ngay trên thiết bị."),
        ("Thách thức: Ràng buộc tài nguyên", "- RAM/Flash cực hạn (tính bằng KB).\n- CPU tốc độ thấp (MHz).\n- Khoảng cách tài nguyên gấp 100.000 lần so với máy tính huấn luyện."),
        ("Thách thức: Training-Deployment Gap", "- Sai lệch công thức toán học giữa Python và C++.\n- Thứ tự đặc trưng không đồng nhất.\n- Các nền tảng hiện có là 'Hộp đen', khó kiểm soát."),
        ("Nền tảng hiện có & Hạn chế", "- Edge Impulse: Phí cao, hộp đen.\n- TFLite Micro: Chỉ hỗ trợ Deep Learning.\n- Thiếu giao diện phân đoạn dữ liệu tương tác."),
        ("Mục tiêu nghiên cứu", "- Xây dựng Framework mã nguồn mở.\n- Tự động hóa quy trình End-to-End.\n- Đảm bảo Training-Deployment Parity.\n- Hỗ trợ đa nền tảng phần cứng."),
        ("Đóng góp của luận văn", "- Giao diện tiền xử lý tương tác kéo-thả.\n- Cơ chế xác minh tương đồng 12 điểm.\n- Kiến trúc sinh mã đa chiều (Factory Pattern)."),
        ("Cơ sở lý thuyết: Dữ liệu IMU", "- Gia tốc kế 3 trục (Chuyển động tuyến tính).\n- Con quay hồi chuyển 3 trục (Vận tốc góc).\n- Tần số lấy mẫu: 50-100Hz."),
        ("Cơ sở lý thuyết: Bộ lọc tín hiệu", "- Butterworth LP (Lọc nhiễu).\n- Savitzky-Golay (Bảo tồn đỉnh).\n- Kalman Filter (Giải pháp tối ưu cho parity gap = 0)."),
        ("Cơ sở lý thuyết: Cửa sổ trượt", "- Phân đoạn chuỗi thời gian.\n- Kích thước cửa sổ: 1.5s (150 mẫu).\n- Chồng lấp: 50% (75 mẫu)."),
        ("Phương pháp: Tiền xử lý tương tác", "- Người dùng kéo-thả trên biểu đồ để chọn đoạn dữ liệu sạch.\n- Giảm 60-80% thời gian gán nhãn thủ công.\n- Tách biệt kiểm soát chất lượng và tự động hóa."),
        ("Phương pháp: Đặc trưng bất biến hướng", "- Tính toán trên Magnitude (độ lớn vector).\n- Không bị ảnh hưởng bởi hướng đeo thiết bị.\n- 63 đặc trưng (Thời gian & Tần số)."),
        ("Phương pháp: Thuật toán AI", "- Machine Learning: Random Forest, SVM, MLP.\n- Deep Learning: 1D-CNN, 2D-CNN.\n- Hỗ trợ PyTorch & Scikit-learn."),
        ("Thiết kế hệ thống: Kiến trúc Framework", "- 6 Modules: Dữ liệu -> Tiền xử lý -> Đặc trưng -> Huấn luyện -> Tạo mã -> Kiểm tra."),
        ("Thiết kế hệ thống: Bộ sinh mã (Generator)", "- Factory Pattern điều phối.\n- 3 Trục: Loại mô hình x Nền tảng x Backend.\n- Tự động sắp xếp lại thứ tự đặc trưng cho C++."),
        ("Đảm bảo tính Parity", "- Khớp công thức Kurtosis/Skewness với Pandas.\n- Edge-value replication thay cho zero-padding.\n- Scaler sync: Nhúng tham số chuẩn hóa vào mã nguồn."),
        ("Thiết lập thực nghiệm", "- Phần cứng: Seeed XIAO nRF52840 (Cortex-M4F).\n- Tập dữ liệu: 5 lớp hoạt động (Running, Walking, Still, Up, Down)."),
        ("Kết quả: Hiệu suất huấn luyện", "- Độ chính xác Offline: 97.38% - 99.13%.\n- CNN và MLP đạt kết quả cao nhất (99.13%)."),
        ("Kết quả: Phân tích nhầm lẫn", "- Running/Still: F1 = 1.00.\n- Nhầm lẫn nhỏ ở các lớp biến thể Walking (Up/Down)."),
        ("Kết quả: Triển khai thực tế", "- 1D-CNN hoạt động đáng tin cậy nhất trên thiết bị thực.\n- Thời gian suy luận: ~20ms.\n- Độ chính xác thực tế khớp với tập kiểm tra."),
        ("So sánh đối chứng (Edge Impulse)", "- Framework: 145KB Flash (E.I: 417KB).\n- Framework: F1=1.00 (E.I: 0.96).\n- Framework: Minh bạch 100% mã nguồn."),
        ("Phân tích các chế độ tối ưu hóa", "- Accuracy (Float32).\n- Balanced (3 chữ số thập phân).\n- Power (Int16 + DSP)."),
        ("Kết luận", "- Framework giải quyết tốt bài toán Edge AI cho HAR.\n- Đảm bảo tính minh bạch và tối ưu tài nguyên."),
        ("Hướng phát triển tương lai", "- Hợp nhất cảm biến (IMU + Nhịp tim).\n- Học liên bang (Federated Learning).\n- Tích hợp Transformer cho chuỗi thời gian."),
    ]
 
    for title, content in slides_content:
        add_slide(prs, title, content)
 
    prs.save('Luan_van_Thac_si_Nguyen_Truong_Minh_Hoang.pptx')
    print("Đã tạo file PowerPoint thành công!")
 
if __name__ == "__main__":
    create_presentation()